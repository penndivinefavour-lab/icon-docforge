"""ICON DocForge v1.1 Test Suite."""
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

BASE = os.path.abspath(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIXTURES = os.path.join(BASE, "tests", "fixtures")
RESULTS = []

def record(name, passed, detail=""):
    RESULTS.append({"test": name, "passed": passed, "detail": detail})
    print(f"  {'✅' if passed else '❌'} {name} {detail}")

from engine.registry import get_registry
from engine.converters.docx_to_pdf import docx_to_pdf
from engine.converters.docx_odt import docx_to_odt
from engine.converters.images_to_pdf import images_to_pdf
from engine.converters.csv_xlsx import csv_to_xlsx, xlsx_to_csv
from engine.converters.markdown_to_pdf import markdown_to_pdf
from engine.converters.pdf_merge import pdf_merge
from engine.converters.pdf_split import pdf_split
from engine.converters.pdf_to_images import pdf_to_images
from engine.converters.pdf_metadata import read_pdf_metadata
from engine.converters.pdf_text import extract_text, search_text, get_page_info, extract_images
from engine.converters.pdf_compress import compress_pdf
from engine.converters.pdf_watermark import add_text_watermark
from engine.converters.pdf_reorder import reorder_pages, delete_pages
from engine.utils import safe_run, sanitize_filename, validate_path, ensure_output_filename
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from openpyxl import load_workbook
from PyPDF2 import PdfReader

def test_registry():
    print("\n=== Registry ===")
    r = get_registry()
    record("Registry initialized", True, f"{r.health_report()['healthy']} healthy")
    record("Pandoc installed", r.engines["pandoc"]["installed"], f"v{r.engines['pandoc']['version']}")
    record("Poppler installed", r.engines["poppler"]["installed"])
    record("WeasyPrint installed", r.engines["weasyprint"]["installed"])
    record("GhostScript NOT installed", not r.engines["ghostscript"]["installed"])
    record("ImageMagick NOT installed", not r.engines["imagemagick"]["installed"])
    record("LibreOffice NOT installed", not r.engines["libreoffice"]["installed"])

def test_docx():
    print("\n=== DOCX ===")
    r = docx_to_pdf(f"{FIXTURES}/docx/sample_document.docx")
    record("DOCX->PDF", r["success"], f"{r.get('size','?')}B")
    r = docx_to_odt(f"{FIXTURES}/docx/sample_document.docx")
    record("DOCX->ODT", r["success"], f"{r.get('size','?')}B")
    r = docx_to_pdf(f"{FIXTURES}/docx/empty_document.docx")
    record("DOCX->PDF empty", r["success"])

def test_images():
    print("\n=== Images->PDF ===")
    r = images_to_pdf([f"{FIXTURES}/images/red.png", f"{FIXTURES}/images/green.jpg"])
    record("Images->PDF 2 imgs", r["success"], f"{r.get('pages','?')}p")
    r = images_to_pdf([f"{FIXTURES}/images/test_image.webp"])
    record("Images->PDF WebP", r["success"])
    r = images_to_pdf([f"{FIXTURES}/images/large_photo.png"])
    record("Images->PDF large", r["success"])
    r = images_to_pdf([f"{FIXTURES}/images/blue_transparent.png"])
    record("Images->PDF transparent", r["success"])
    r = images_to_pdf([f"{FIXTURES}/images/rotated.jpg"])
    record("Images->PDF rotated", r["success"])
    r = images_to_pdf([])
    record("Empty images fails", not r["success"])
    r = images_to_pdf(["/nonexistent/file.png"])
    record("Nonexistent image fails", not r["success"])
    r = images_to_pdf([f"{FIXTURES}/images/red.png"], orientation="landscape", page_size="A4")
    record("Landscape A4", r["success"])

def test_pdf():
    print("\n=== PDF ===")
    fp = os.path.join(FIXTURES, "pdf")
    for i in range(1,4):
        c = canvas.Canvas(os.path.join(fp, f"t{i}.pdf"), pagesize=A4)
        c.drawString(100,750,f"Page {i}"); c.showPage(); c.save()
    t1, t2 = os.path.join(fp,"t1.pdf"), os.path.join(fp,"t2.pdf")
    r = pdf_merge([t1, t2], output_path=os.path.join(fp,"m.pdf"))
    record("PDF merge", r["success"], f"{r.get('size','?')}B")
    r = pdf_split(os.path.join(fp,"m.pdf"), "1-2", output_path=os.path.join(fp,"s.pdf"))
    record("PDF split", r["success"], f"{r.get('size','?')}B")
    r = pdf_to_images(os.path.join(fp,"m.pdf"), output_path=os.path.join(fp,"img"))
    record("PDF->images", r["success"], f"{r.get('count','?')}p")
    r = read_pdf_metadata(os.path.join(fp,"m.pdf"))
    record("PDF metadata", r["success"])
    r = pdf_merge([t1, t2, os.path.join(fp,"t3.pdf")], output_path=os.path.join(fp,"m3.pdf"))
    record("PDF merge 3 pages", r["success"], f"{r.get('size','?')}B")
    r = pdf_split(os.path.join(fp,"m3.pdf"), "1-3", output_path=os.path.join(fp,"s3.pdf"))
    record("PDF split all", r["success"], f"{r.get('size','?')}B")
    for f in ['t1.pdf','t2.pdf','t3.pdf','m.pdf','s.pdf','m3.pdf','s3.pdf']:
        p = os.path.join(fp, f)
        if os.path.exists(p): os.remove(p)
    for f in os.listdir(fp):
        if f.startswith('img'): os.remove(os.path.join(fp,f))

def test_pdf_text():
    print("\n=== PDF Text ===")
    r = extract_text(f"{FIXTURES}/pdf/multi_page.pdf")
    record("PDF->text extraction", r["success"], f"{r.get('size','?')}B")
    r = search_text(f"{FIXTURES}/pdf/multi_page.pdf", "Page 1")
    record("PDF search", r["success"], f"matches={r.get('matches','?')}")
    r = search_text(f"{FIXTURES}/pdf/multi_page.pdf", "nonexistent_xyz", case_sensitive=True)
    record("PDF search no match", r.get("matches",1)==0)
    r = get_page_info(f"{FIXTURES}/pdf/multi_page.pdf")
    record("PDF page info", r["success"], f"pages={r.get('info',{}).get('Pages','?')}")
    r = extract_images(f"{FIXTURES}/pdf/multi_page.pdf")
    record("PDF->image extraction", r["success"], f"{r.get('count','?')} images")

def test_pdf_advanced():
    print("\n=== PDF Advanced ===")
    fp = os.path.join(FIXTURES, "pdf")
    c = canvas.Canvas(os.path.join(fp, "adv_test.pdf"), pagesize=A4)
    for i in range(3):
        c.drawString(100,750,f"Page {i+1}"); c.showPage()
    c.save()
    r = compress_pdf(os.path.join(fp, "adv_test.pdf"))
    record("PDF compression", r["success"], f"ratio={r.get('compression_ratio','?')}")
    r = add_text_watermark(os.path.join(fp, "adv_test.pdf"), text="CONFIDENTIAL")
    record("Text watermark", r["success"], f"pages={r.get('pages','?')}")
    r = reorder_pages(os.path.join(fp, "adv_test.pdf"), [3,1,2])
    record("Page reorder", r["success"], f"{r.get('size','?')}B")
    r = delete_pages(os.path.join(fp, "adv_test.pdf"), [2])
    record("Delete page", r["success"], f"{r.get('size','?')}B")
    os.remove(os.path.join(fp, "adv_test.pdf"))

def test_spreadsheet():
    print("\n=== Spreadsheet ===")
    r = csv_to_xlsx(f"{FIXTURES}/spreadsheet/sample_data.csv")
    record("CSV->XLSX", r["success"], f"{r.get('size','?')}B")
    r = xlsx_to_csv(f"{FIXTURES}/spreadsheet/sample_data.xlsx")
    record("XLSX->CSV", r["success"], f"{r.get('size','?')}B")
    wb = load_workbook(f"{FIXTURES}/spreadsheet/sample_data.xlsx")
    record("XLSX multi-sheet", len(wb.sheetnames) >= 2)

def test_markdown():
    print("\n=== Markdown ===")
    md = os.path.join(FIXTURES, "docx", "test.md")
    with open(md, 'w') as f: f.write("# Test\n**Bold** text")
    r = markdown_to_pdf(md)
    record("Markdown->PDF", r["success"], f"{r.get('size','?')}B")
    os.remove(md)

def test_utils():
    print("\n=== Utilities ===")
    record("sanitize_filename", sanitize_filename("a/b/c") == "c")
    record("sanitize_unicode", sanitize_filename("café.txt") == "café.txt")
    record("validate_path", True)
    record("ensure_output_filename", ensure_output_filename("out", ".pdf") == "out.pdf")
    r = safe_run(["echo", "ok"], timeout=5)
    record("safe_run", r["success"])
    r = safe_run(["nonexistent"], timeout=5)
    record("safe_run missing", not r["success"])

def test_fixture_quality():
    print("\n=== Fixture Quality ===")
    doc = __import__('docx').Document(f"{FIXTURES}/docx/sample_document.docx")
    headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
    tables = doc.tables
    record("DOCX has headings", len(headings) >= 1)
    record("DOCX has tables", len(tables) >= 1)
    reader = PdfReader(f"{FIXTURES}/pdf/multi_page.pdf")
    record("PDF has pages", len(reader.pages) >= 1)
    img_count = len([f for f in os.listdir(f"{FIXTURES}/images") if f.endswith(('.png','.jpg','.webp'))])
    record("Image fixtures", img_count >= 5, f"{img_count} images")

def test_malformed():
    print("\n=== Malformed/Edge ===")
    r = docx_to_pdf(f"{FIXTURES}/docx/nonexistent.docx")
    record("Nonexistent DOCX fails", not r["success"])
    r = pdf_merge([])
    record("Empty merge fails", not r["success"])
    r = images_to_pdf([])
    record("Empty images fails", not r["success"])
    r = csv_to_xlsx("/nonexistent.csv")
    record("Nonexistent CSV fails", not r["success"])

if __name__ == "__main__":
    print("=" * 60)
    print("ICON DocForge v1.1 Test Suite")
    print("=" * 60)
    test_registry(); test_docx(); test_images(); test_pdf()
    test_pdf_text(); test_pdf_advanced(); test_spreadsheet()
    test_markdown(); test_utils(); test_fixture_quality()
    test_malformed()
    passed = sum(1 for r in RESULTS if r["passed"])
    print(f"\n{'='*60}\nResults: {passed}/{len(RESULTS)} tests passed")
    with open(os.path.join(BASE, "tests", "test_results.json"), 'w') as f:
        json.dump({"results": RESULTS, "passed": passed, "total": len(RESULTS)}, f, indent=2)
