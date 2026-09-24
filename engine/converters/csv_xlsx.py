"""CSV ↔ XLSX converter using openpyxl."""
import os
import csv
from openpyxl import Workbook, load_workbook
from engine.config import OUTPUT_DIR


def csv_to_xlsx(input_path: str, output_path: str = None) -> dict:
    """Convert CSV to XLSX using openpyxl."""
    if not os.path.exists(input_path):
        return {"success": False, "error": f"CSV not found"}
    
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, os.path.splitext(os.path.basename(input_path))[0] + ".xlsx")
    
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        with open(input_path, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                ws.append(row)
        
        wb.save(output_path)
        return {"success": True, "output": output_path, "size": os.path.getsize(output_path), "method": "openpyxl"}
    except Exception as e:
        return {"success": False, "error": str(e)}


def xlsx_to_csv(input_path: str, output_path: str = None) -> dict:
    """Convert XLSX to CSV using openpyxl."""
    if not os.path.exists(input_path):
        return {"success": False, "error": f"XLSX not found"}
    
    if output_path is None:
        output_path = os.path.join(OUTPUT_DIR, os.path.splitext(os.path.basename(input_path))[0] + ".csv")
    
    try:
        wb = load_workbook(input_path, read_only=True)
        ws = wb.active
        
        with open(output_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(row)
        
        wb.close()
        return {"success": True, "output": output_path, "size": os.path.getsize(output_path), "method": "openpyxl"}
    except Exception as e:
        return {"success": False, "error": str(e)}
